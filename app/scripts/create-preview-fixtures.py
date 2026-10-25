from pathlib import Path
from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pptx import Presentation

root = Path(__file__).resolve().parents[2] / 'tmp' / 'preview-fixtures'
root.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_heading('Print layout verification', 0)
doc.add_paragraph('This document keeps its headings, spacing, table, and page breaks.')
table = doc.add_table(rows=2, cols=2)
table.style = 'Table Grid'
table.cell(0, 0).text = 'Name'
table.cell(0, 1).text = 'Copies'
table.cell(1, 0).text = 'Original customer document'
table.cell(1, 1).text = '2'
image = root.parent / 'editor-fixtures' / 'sample.png'
if image.exists():
    doc.add_picture(str(image), width=Inches(2))
doc.add_page_break()
doc.add_heading('Second page', 1)
doc.add_paragraph('This is page two of the actual document.')
doc.save(root / 'layout.docx')

wb = Workbook()
sheet = wb.active
sheet.title = 'Print sheet'
sheet.append(['Item', 'Quantity', 'Price'])
sheet.append(['Customer print job', 2, 5])
sheet.append(['Total', '=B2', '=B2*C2'])
for cell in sheet[1]:
    cell.font = Font(bold=True, color='FFFFFF')
    cell.fill = PatternFill('solid', fgColor='17304D')
sheet.column_dimensions['A'].width = 30
sheet.column_dimensions['B'].width = 18
sheet.column_dimensions['C'].width = 18
sheet.print_area = 'A1:C3'
sheet.sheet_properties.pageSetUpPr.fitToPage = True
sheet.page_setup.orientation = 'landscape'
sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
sheet.page_setup.fitToWidth = 1
sheet.page_setup.fitToHeight = 1
wb.save(root / 'layout.xlsx')

slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text = 'Original slide layout'
slide.placeholders[1].text = 'Slide text and shapes stay in position.'
slides.save(root / 'layout.pptx')
print(root)

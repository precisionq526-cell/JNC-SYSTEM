"""Create an isolated workbook for native Excel edit/print verification."""
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, Protection
from openpyxl.drawing.image import Image
from PIL import Image as PILImage, ImageDraw

folder = Path(__file__).resolve().parents[2] / 'tmp' / 'editor-header-check'
folder.mkdir(parents=True, exist_ok=True)
logo = PILImage.new('RGB', (260, 70), '#14263f')
ImageDraw.Draw(logo).text((15, 25), 'WORKSHEET LAYOUT CHECK', fill='white')
logo.save(folder / 'worksheet-logo.png')
book = Workbook()
sheet = book.active
sheet.title = 'Print job'
sheet.merge_cells('A1:E1')
sheet['A1'] = 'Worksheet editor verification'
sheet['A1'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
sheet['A1'].fill = PatternFill('solid', fgColor='14263F')
sheet['A1'].alignment = Alignment(horizontal='center')
sheet.row_dimensions[1].height = 36
for column, width in [('A', 30), ('B', 14), ('C', 18), ('D', 18), ('E', 22)]:
    sheet.column_dimensions[column].width = width
sheet.append(['Item', 'Quantity', 'Unit price', 'Line total', 'Date'])
sheet.append(['Document printing', 2, 25, '=B3*C3', datetime(2026, 10, 20)])
sheet.append(['Binding', 1, 40, '=B4*C4', datetime(2026, 10, 20)])
sheet['A5'] = 'Additional item'
sheet['D5'] = '=B5*C5'
sheet['A7'] = 'Total'
sheet['D7'] = '=SUM(D3:D5)'
for row in sheet.iter_rows(min_row=2, max_row=7, max_col=5):
    for cell in row:
        cell.font = Font(name='Calibri', size=11, bold=cell.row in (2, 7))
        cell.border = Border(bottom=Side(style='thin', color='D3DAE3'))
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        if cell.column in (3, 4): cell.number_format = '"PHP "#,##0.00'
        if cell.column == 5: cell.number_format = 'yyyy-mm-dd'
for r in range(2, 8): sheet.row_dimensions[r].height = 25
sheet.add_image(Image(folder / 'worksheet-logo.png'), 'A10')
sheet.print_area = 'A1:E16'
sheet.page_setup.orientation = 'landscape'
sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
sheet.page_setup.fitToWidth = 1
sheet.page_setup.fitToHeight = 1
sheet.sheet_properties.pageSetUpPr.fitToPage = True
sheet.oddHeader.center.text = 'Original print header'
sheet.oddFooter.center.text = 'Page &P of &N'
reference = book.create_sheet('Reference')
reference.append(['This sheet must stay unchanged', 123])
reference.print_area = 'A1:B2'
reference.column_dimensions['A'].width = 42
protected = book.create_sheet('Protected')
protected['A1'] = 'Locked cell'
protected['B1'] = 'Unlocked cell'
protected['B1'].protection = Protection(locked=False)
protected.protection.sheet = True
protected.sheet_state = 'hidden'
book.save(folder / 'worksheet-original.xlsx')
print('Created workbook fixture with merged title, styled cells, formulas, dates, print area, picture and protected sheet.')
